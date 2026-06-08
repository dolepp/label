"""Admin-generated XLSX report sending callback."""
from __future__ import annotations

import logging
import os
import tempfile
from datetime import datetime
from typing import Any

from telebot import types

from db.repositories.admin_reports import get_report_xlsx_payload, mark_report_sent

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    XLSX_AVAILABLE = True
except ImportError:
    openpyxl = None
    Font = Alignment = PatternFill = Border = Side = None
    XLSX_AVAILABLE = False

logger = logging.getLogger(__name__)

bot = None


def configure_admin_send_reports(**context: Any) -> None:
    globals().update({key: value for key, value in context.items() if value is not None})


def _require(name: str) -> Any:
    value = globals().get(name)
    if value is None:
        raise RuntimeError(f"admin send reports dependency is not configured: {name}")
    return value


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _autosize_columns(ws, max_width: int = 50) -> None:
    for column in ws.columns:
        max_length = 0
        column_letter = _column_letter(column[0].column)
        for cell in column:
            try:
                max_length = max(max_length, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, max_width)

def create_detailed_xlsx_report(user_data, releases_data=None, promo_codes_data=None, orders_data=None):
    """Создание детального XLSX отчета с несколькими листами"""
    if not XLSX_AVAILABLE:
        raise ValueError("Библиотека openpyxl не установлена. Обратитесь к администратору.")

    # Создаем новую книгу Excel
    wb = openpyxl.Workbook()

    # Удаляем лист по умолчанию
    wb.remove(wb.active)

    # Стили
    header_font = Font(bold=True, size=14, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    subheader_font = Font(bold=True, size=12, color="000000")
    subheader_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Лист 1: Информация о пользователе
    ws_user = wb.create_sheet("Пользователь")
    ws_user['A1'] = f"ОТЧЕТ ПО ПОЛЬЗОВАТЕЛЮ: {user_data.get('name', 'Неизвестно')}"
    ws_user['A1'].font = header_font
    ws_user['A1'].fill = header_fill
    ws_user['A1'].alignment = header_alignment
    ws_user.merge_cells('A1:H1')

    user_info = [
        ["Имя:", user_data.get('name', 'Не указано')],
        ["Telegram ID:", str(user_data.get('telegram_id', 'Не указано'))],
        ["Username:", user_data.get('tg', 'Не указано')],
        ["Email:", user_data.get('email', 'Не указано')],
        ["Дата регистрации:", str(user_data.get('created_at', 'Не указано'))],
        ["Статус:", user_data.get('status', 'Не указано')],
        ["Роль:", user_data.get('role', 'Не указано')]
    ]

    for i, (label, value) in enumerate(user_info, start=3):
        ws_user[f'A{i}'] = label
        ws_user[f'B{i}'] = value
        ws_user[f'A{i}'].font = Font(bold=True)
        ws_user[f'A{i}'].border = thin_border
        ws_user[f'B{i}'].border = thin_border

    # Лист 2: Релизы
    if releases_data:
        ws_releases = wb.create_sheet("Релизы")
        ws_releases['A1'] = "РЕЛИЗЫ ПОЛЬЗОВАТЕЛЯ"
        ws_releases['A1'].font = header_font
        ws_releases['A1'].fill = header_fill
        ws_releases['A1'].alignment = header_alignment
        ws_releases.merge_cells('A1:H1')

        release_headers = [
            "ID", "Название", "Тип", "Статус", "Дата создания",
            "Дата обновления", "Количество треков", "Описание"
        ]

        for col, header in enumerate(release_headers, start=1):
            cell = ws_releases.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for row_idx, release in enumerate(releases_data, start=4):
            for col_idx, value in enumerate([
                release.get('id', ''),
                release.get('name', ''),
                release.get('type', ''),
                release.get('status', ''),
                str(release.get('created_at', '')),
                str(release.get('updated_at', '')),
                release.get('track_count', 0),
                release.get('description', '')
            ], start=1):
                cell = ws_releases.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Лист 3: Промокоды (если есть)
    if promo_codes_data:
        ws_promo = wb.create_sheet("Промокоды")
        ws_promo['A1'] = "ПРОМОКОДЫ ПОЛЬЗОВАТЕЛЯ"
        ws_promo['A1'].font = header_font
        ws_promo['A1'].fill = header_fill
        ws_promo['A1'].alignment = header_alignment
        ws_promo.merge_cells('A1:F1')

        promo_headers = [
            "ID", "Код", "Сумма", "Макс. использований", "Текущие использования",
            "Дата истечения", "Статус"
        ]

        for col, header in enumerate(promo_headers, start=1):
            cell = ws_promo.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for row_idx, promo in enumerate(promo_codes_data, start=4):
            for col_idx, value in enumerate([
                promo.get('id', ''),
                promo.get('code', ''),
                promo.get('amount', ''),
                promo.get('max_uses', 'Без ограничений'),
                promo.get('current_uses', 0),
                str(promo.get('expires_at', 'Без срока')),
                'Активен' if promo.get('is_active') else 'Неактивен'
            ], start=1):
                cell = ws_promo.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Лист 4: Заказы (если есть)
    if orders_data:
        ws_orders = wb.create_sheet("Заказы")
        ws_orders['A1'] = "ЗАКАЗЫ ПОЛЬЗОВАТЕЛЯ"
        ws_orders['A1'].font = header_font
        ws_orders['A1'].fill = header_fill
        ws_orders['A1'].alignment = header_alignment
        ws_orders.merge_cells('A1:G1')

        order_headers = [
            "ID", "Тип услуги", "Статус", "Сумма", "Дата создания",
            "Дата завершения", "Описание"
        ]

        for col, header in enumerate(order_headers, start=1):
            cell = ws_orders.cell(row=3, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for row_idx, order in enumerate(orders_data, start=4):
            for col_idx, value in enumerate([
                order.get('id', ''),
                order.get('service_type', ''),
                order.get('status', ''),
                order.get('amount', ''),
                str(order.get('created_at', '')),
                str(order.get('completed_at', '')),
                order.get('description', '')
            ], start=1):
                cell = ws_orders.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="top")

    # Лист 5: Сводка
    ws_summary = wb.create_sheet("Сводка")
    ws_summary['A1'] = "СВОДКА ПО ПОЛЬЗОВАТЕЛЮ"
    ws_summary['A1'].font = header_font
    ws_summary['A1'].fill = header_fill
    ws_summary['A1'].alignment = header_alignment
    ws_summary.merge_cells('A1:D1')

    summary_data = [
        ["Общее количество релизов:", len(releases_data) if releases_data else 0],
        ["Активных релизов:", len([r for r in (releases_data or []) if r.get('status') == 'active'] or 0)],
        ["Промокодов:", len(promo_codes_data) if promo_codes_data else 0],
        ["Заказов:", len(orders_data) if orders_data else 0],
        ["Дата генерации отчета:", datetime.now().strftime("%d.%m.%Y %H:%M:%S")]
    ]

    for i, (label, value) in enumerate(summary_data, start=3):
        ws_summary[f'A{i}'] = label
        ws_summary[f'B{i}'] = value
        ws_summary[f'A{i}'].font = Font(bold=True)
        ws_summary[f'A{i}'].border = thin_border
        ws_summary[f'B{i}'].border = thin_border

    # Автоматическая ширина столбцов для всех листов
    for ws in wb.worksheets:
        for column in ws.columns:
            max_length = 0
            column_letter = _column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    return wb

def send_xlsx_report(chat_id, report_data, filename="report.xlsx"):
    """Отправка XLSX отчета пользователю"""
    tmp_file_path = None
    try:
        # Сохраняем отчет во временный файл
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            report_data.save(tmp_file.name)
            tmp_file_path = tmp_file.name

        # Отправляем файл
        with open(tmp_file_path, 'rb') as file:
            _require("bot").send_document(
                chat_id,
                file,
                caption="📊 Ваш отчет в формате XLSX готов!\n\n"
                        "📋 Отчет содержит детальную информацию по всем разделам.\n"
                        "💾 Файл сохранен в формате Excel для удобного просмотра и анализа.",
                visible_file_name=filename
            )

        return True

    except Exception as e:
        logger.error(f"Ошибка при отправке XLSX отчета: {e}")
        _require("bot").send_message(chat_id, "❌ Ошибка при отправке отчета. Попробуйте еще раз.")
        return False
    finally:
        if tmp_file_path and os.path.exists(tmp_file_path):
            try:
                os.unlink(tmp_file_path)
            except OSError:
                logger.warning("Could not remove temporary XLSX report: %s", tmp_file_path)

def handle_admin_send_report(call):
    """Handle admin sending completed report to user"""
    report_id = int(call.data.split("_")[3])
    current_bot = _require("bot")

    try:
        payload = get_report_xlsx_payload(report_id)
        if payload is None:
            current_bot.answer_callback_query(call.id, "❌ Отчет не найден")
            return
        if payload.get("status") == "missing_user":
            current_bot.answer_callback_query(call.id, "❌ Данные пользователя не найдены")
            return

        if not XLSX_AVAILABLE:
            current_bot.answer_callback_query(call.id, "❌ Модуль Excel недоступен")
            return

        report = payload["report"]
        wb = create_detailed_xlsx_report(
            payload["user"],
            payload["releases"],
            payload["promo_codes"],
            payload["orders"],
        )

        marked = mark_report_sent(report_id, call.from_user.id)
        if marked is None:
            current_bot.answer_callback_query(call.id, "❌ Ошибка подключения к базе данных")
            return
        if not marked:
            current_bot.answer_callback_query(call.id, "❌ Отчет не найден")
            return

        try:
            user_name = report["user_name"]
            username = report.get("username")
            request_type = report["request_type"]
            filename = f"Отчет_{user_name}_{datetime.now().strftime('%d%m%Y')}.xlsx"
            if send_xlsx_report(report["user_id"], wb, filename):
                current_bot.edit_message_text(
                    f"✅ Отчет #{report_id} успешно отправлен пользователю {user_name} (@{username or 'без username'})\n\n"
                    f"📊 Тип отчета: {request_type}\n"
                    f"📅 Дата отправки: {datetime.now().strftime('%d.%m.%Y %H:%M')}\n\n"
                    f"📎 Отчет отправлен в формате XLSX с детальной информацией\n"
                    f"📋 Содержит {len(wb.worksheets)} листов с данными",
                    call.message.chat.id,
                    call.message.message_id,
                    reply_markup=types.InlineKeyboardMarkup().add(
                        types.InlineKeyboardButton("◀️ Назад к запросам", callback_data="admin_report_requests")
                    )
                )
            else:
                current_bot.answer_callback_query(call.id, "❌ Ошибка при отправке отчета")

        except Exception as e:
            logger.error(f"Error sending report to user: {e}")
            current_bot.answer_callback_query(call.id, f"❌ Ошибка отправки отчета: {str(e)}")

    except Exception as e:
        logger.error(f"Error in handle_admin_send_report: {e}")
        current_bot.answer_callback_query(call.id, f"❌ Ошибка: {str(e)}")


def register_admin_send_report_handlers(bot, context: dict | None = None) -> None:
    configure_admin_send_reports(bot=bot, **(context or {}))

    @bot.callback_query_handler(func=lambda call: (getattr(call, "data", "") or "").startswith("admin_send_report_"))
    def admin_send_report_callback(call):
        handle_admin_send_report(call)
