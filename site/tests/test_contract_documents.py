import sys
import unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from contract_documents import render_contract, normalized_profile

PROFILE = dict(full_name='Тестов Тест Тестович', passport='1111 222222',passport_issued='Тестовый отдел',issue_date='2020-01-01',department_code='111-222',birth_date='1990-01-01',birth_place='Тестовый город',address='Тестовая улица, 1')
class ContractTests(unittest.TestCase):
    def test_album_preserves_licensee_and_terms(self):
        out,number=render_contract(PROFILE,dict(id=42,release_name='Тестовый альбом',artist_name='Тест',release_date='2026-12-01'),[dict(release_name='Первый',music_author='Автор 1'),dict(release_name='Второй',text_author='Автор 2')])
        doc=Document(out);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('Чабин Илья Анатольевич',text)
        self.assertIn('(80) % (восемьдесят процентов)',text)
        self.assertIn('Тестов Тест Тестович',text)
        self.assertNotIn('{{',text)
        self.assertNotIn('Фамилия Имя Отчество',text)
        self.assertNotIn('1234 567890',text)
        self.assertEqual(len(doc.tables[1].rows),3)
        self.assertIn('Первый',doc.tables[1].rows[1].cells[1].text)
        self.assertIn('Второй',doc.tables[1].rows[2].cells[1].text)
        self.assertIn(number,text)
        self.assertIn('40817810000035696053',doc.tables[0].rows[1].cells[1].text)
        self.assertNotIn('40817810000035696053',doc.tables[0].rows[1].cells[0].text)
    def test_required_and_invalid_date(self):
        with self.assertRaises(ValueError): normalized_profile({})
        with self.assertRaises(ValueError): normalized_profile(dict(PROFILE,birth_date='2099-01-01'))
        with self.assertRaises(ValueError): normalized_profile(dict(PROFILE,passport='123'))
if __name__=='__main__': unittest.main()
