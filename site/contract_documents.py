"""Fill the owner's supplied DOCX without replacing legal clauses or styles."""
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches

TEMPLATE_PATH = Path(__file__).parent / 'templates/license_agreement_v1.docx'
TEMPLATE_VERSION = '2026-10-06-v1'
FIELDS = {
    'full_name': 'ФИО', 'nickname': 'Псевдоним', 'passport': 'Серия и номер паспорта',
    'passport_issued': 'Кем выдан паспорт', 'issue_date': 'Дата выдачи',
    'department_code': 'Код подразделения', 'birth_date': 'Дата рождения',
    'birth_place': 'Место рождения', 'address': 'Адрес регистрации', 'snils': 'СНИЛС', 'inn': 'ИНН',
    'bank_name': 'Название банка', 'bank_account': 'Расчётный / лицевой счёт',
    'bank_correspondent_account': 'Корреспондентский счёт', 'bank_bic': 'БИК',
    'bank_inn': 'ИНН банка', 'bank_kpp': 'КПП банка', 'phonogram_producer': 'Изготовитель фонограммы',
}
REQUIRED = ('full_name', 'passport', 'passport_issued', 'issue_date', 'department_code', 'birth_date', 'birth_place', 'address')
MONTHS = ('января', 'февраля', 'марта', 'апреля', 'мая', 'июня', 'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря')


def date_value(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    for fmt in ('%Y-%m-%d', '%d.%m.%Y'):
        try: return datetime.strptime(str(value), fmt).date()
        except ValueError: pass
    raise ValueError('Укажите дату в формате ДД.ММ.ГГГГ или ГГГГ-ММ-ДД.')


def normalized_profile(data, partial=False):
    result = {}
    for key in FIELDS:
        value = data.get(key, '')
        if not isinstance(value, str): raise ValueError(f'{FIELDS[key]}: ожидается текст.')
        value = value.strip()
        if len(value) > 600 or any(ord(ch) < 32 and ch not in '\n\t' for ch in value):
            raise ValueError(f'{FIELDS[key]}: слишком длинное или некорректное значение.')
        result[key] = value
    for key in ('issue_date', 'birth_date'):
        if result[key]:
            parsed = date_value(result[key])
            if parsed > date.today(): raise ValueError(f'{FIELDS[key]} не может быть в будущем.')
            result[key] = parsed.isoformat()
    for key, digits in [('passport', 10), ('department_code', 6), ('snils', 11), ('inn', 12), ('bank_account', 20), ('bank_correspondent_account', 20), ('bank_bic', 9), ('bank_inn', 10), ('bank_kpp', 9)]:
        if result[key] and not re.fullmatch(r'[\d\s-]+', result[key]):
            raise ValueError(f'{FIELDS[key]}: используйте цифры, пробелы и дефисы.')
        if result[key] and len(re.sub(r'\D', '', result[key])) != digits:
            raise ValueError(f'{FIELDS[key]}: должно быть {digits} цифр.')
    if not partial:
        missing = [FIELDS[key] for key in REQUIRED if not result[key]]
        if missing: raise ValueError('Заполните реквизиты: ' + ', '.join(missing))
    return result


def replace_text(paragraph_element, replacements):
    """Replace text across DOCX runs, retaining the surrounding run formatting."""
    nodes = paragraph_element.xpath('.//w:t')
    text = ''.join(node.text or '' for node in nodes)
    count = 0
    # A single pass prevents a user value from becoming another placeholder.
    pattern = re.compile('|'.join(re.escape(k) for k in sorted(replacements, key=len, reverse=True))) if replacements else None
    if pattern is None: return 0
    matches = list(pattern.finditer(text))
    offsets = []
    at = 0
    for node in nodes:
        end = at + len(node.text or '')
        offsets.append((node, at, end))
        at = end
    for match in reversed(matches):
        start, end = match.span()
        for node, node_start, node_end in offsets:
            if node_end <= start or node_start >= end: continue
            original = node.text or ''
            left, right = max(0, start - node_start), min(len(original), end - node_start)
            value = str(replacements[match.group()]) if node_start <= start < node_end else ''
            node.text = original[:left] + value + original[right:]
            if node.text and (node.text[0].isspace() or node.text[-1].isspace()):
                node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        count += 1
    return count


def document_roots(document):
    roots = [document.element]
    seen = set()
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            name = str(part.part.partname)
            if name not in seen:
                roots.append(part._element); seen.add(name)
    return roots


def prepare_template(original, target=TEMPLATE_PATH):
    document = Document(original)
    common = {
        'Фамилия Имя Отчество': '{{full_name}}', '1234 567890': '{{passport}}',
        'ПСЕВДОНИМ': '{{nickname}}', 'Фамилия И. О.': '{{initials}}',
        'X/-XX': '{{contract_number}}', '3/-01': '{{contract_number}}',
        '01 января 2026 года': '{{contract_date}}', 'КЕМ ВЫДАН': '{{passport_issued}}',
        'Дата выдачи: 01.01.2001': 'Дата выдачи: {{issue_date}}',
        'Дата рождения: 01.01.2001': 'Дата рождения: {{birth_date}}',
        'Код подразделения: 600-006': 'Код подразделения: {{department_code}}',
        'Место рождения: г. Москва': 'Место рождения: {{birth_place}}',
        'Адрес регистрации: Адрес регистрации': 'Адрес регистрации: {{address}}',
        '123-456-789 01': '{{snils}}', '781432831090': '{{inn}}',
    }
    for root in document_roots(document):
        for paragraph in root.xpath('.//w:p'): replace_text(paragraph, common)
    # Only licensor cells: the licensee's actual bank details must stay intact.
    bank = {'Псевдоним': '{{nickname}}', 'АО ‘T-Банк’': '{{bank_name}}',
            '40817810000035696053': '{{bank_account}}', '30101810145250000974': '{{bank_correspondent_account}}',
            '7710140679': '{{bank_inn}}', '044525974': '{{bank_bic}}', '771301001': '{{bank_kpp}}'}
    for table in document.tables:
        for row in table.rows:
            if not row.cells: continue
            cell = row.cells[0]
            if '{{full_name}}' in cell.text:
                for paragraph in cell._tc.xpath('.//w:p'): replace_text(paragraph, bank)
    track_table = document.tables[1]
    replacements = ['release_name', 'track_name', 'music_author', 'text_author', 'performer', 'phonogram_producer']
    for i, key in enumerate(replacements):
        cell = track_table.rows[1].cells[i]
        for paragraph in cell._tc.xpath('.//w:p'):
            text = ''.join(t.text or '' for t in paragraph.xpath('.//w:t')).strip()
            if text: replace_text(paragraph, {text: '{{' + key + '}}'})
    for paragraph in track_table.rows[1].cells[7]._tc.xpath('.//w:p'): replace_text(paragraph, {'2025': '{{delivery_date}}'})
    # Replace the sample cover in Appendix 2, keeping its paragraph position.
    for paragraph in document.paragraphs:
        if paragraph._p.xpath('.//w:drawing'):
            for child in list(paragraph._p):
                if child.tag.endswith('}r'): paragraph._p.remove(child)
            paragraph.add_run('{{cover}}')
    Path(target).parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def render_contract(profile, release, tracks=None, number=None, cover_path=None, contract_date=None, strict=True):
    values = normalized_profile(profile, partial=not strict)
    today = date_value(contract_date) if contract_date else date.today()
    contract_date_text = f'{today.day:02d} {MONTHS[today.month - 1]} {today.year} года'
    number = number or f'TWAS-{today:%Y%m%d}-{release.get("id", "draft")}'
    names = values['full_name'].split()
    initials = names[0] + ' ' + ' '.join(name[0] + '.' for name in names[1:]) if names else 'Не указано'
    values.update(contract_number=number, contract_date=contract_date_text, initials=initials,
                  nickname=values['nickname'] or release.get('artist_name') or 'Не указано')
    for key in ('issue_date', 'birth_date'):
        if values[key]: values[key] = date_value(values[key]).strftime('%d.%m.%Y')
    document = Document(TEMPLATE_PATH)
    table = document.tables[1]
    source_row = deepcopy(table.rows[1]._tr)
    table._tbl.remove(table.rows[1]._tr)
    tracks = tracks or [release]
    for track in tracks:
        row_element = deepcopy(source_row)
        table._tbl.append(row_element)
        row_values = dict(values, release_name=release.get('release_name') or '',
                          track_name=track.get('release_name') or track.get('track_name') or '',
                          music_author=track.get('music_author') or '', text_author=track.get('text_author') or '',
                          performer=track.get('performer_name') or track.get('artist_name') or release.get('artist_name') or '',
                          phonogram_producer=track.get('phonogram_producer') or values['phonogram_producer'] or values['full_name'],
                          delivery_date=date_value(release['release_date']).strftime('%d.%m.%Y') if release.get('release_date') else today.strftime('%d.%m.%Y'))
        replace = {'{{' + k + '}}': v or 'Не указано' for k, v in row_values.items()}
        for paragraph in row_element.xpath('.//w:p'): replace_text(paragraph, replace)
    for root in document_roots(document):
        for paragraph in root.xpath('.//w:p'):
            replace_text(paragraph, {'{{' + k + '}}': v or 'Не указано' for k, v in values.items()})
    for paragraph in document.paragraphs:
        if '{{cover}}' in paragraph.text:
            replace_text(paragraph._p, {'{{cover}}': '' if cover_path and Path(cover_path).is_file() else 'Обложка релиза не загружена.'})
            if cover_path and Path(cover_path).is_file(): paragraph.add_run().add_picture(str(cover_path), width=Inches(3.8))
    for root in document_roots(document):
        if re.search(r'\{\{[a-z_]+\}\}', ''.join(root.xpath('.//w:t/text()'))):
            raise ValueError('В шаблоне остались незаполненные поля.')
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output, number
